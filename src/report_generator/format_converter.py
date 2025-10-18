"""
Report format converter for CHIPS Act entity tracking.
Converts markdown reports to Word (.docx) and text (.txt) formats.
"""

import re
from pathlib import Path
from typing import Optional
from docx import Document
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.shared import OxmlElement, qn
from loguru import logger


class ReportFormatConverter:
    """Convert markdown reports to different formats."""
    
    def __init__(self):
        self.supported_formats = ['md', 'docx', 'txt']
    
    def convert_to_word(self, markdown_content: str, output_path: str) -> str:
        """Convert markdown content to Word document."""
        try:
            doc = Document()
            
            # Set document margins
            sections = doc.sections
            for section in sections:
                section.top_margin = Inches(1)
                section.bottom_margin = Inches(1)
                section.left_margin = Inches(1)
                section.right_margin = Inches(1)
            
            # Parse markdown content
            lines = markdown_content.split('\n')
            i = 0
            
            while i < len(lines):
                line = lines[i].strip()
                
                if not line:
                    i += 1
                    continue
                
                # Handle headers
                if line.startswith('#'):
                    level = len(line) - len(line.lstrip('#'))
                    header_text = line.lstrip('# ').strip()
                    
                    if level == 1:
                        heading = doc.add_heading(header_text, level=1)
                        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    elif level == 2:
                        doc.add_heading(header_text, level=2)
                    elif level == 3:
                        doc.add_heading(header_text, level=3)
                    else:
                        doc.add_heading(header_text, level=4)
                
                # Handle bullet points
                elif line.startswith('- ') or line.startswith('* '):
                    bullet_text = line[2:].strip()
                    doc.add_paragraph(bullet_text, style='List Bullet')
                
                # Handle numbered lists
                elif re.match(r'^\d+\.', line):
                    numbered_text = re.sub(r'^\d+\.\s*', '', line)
                    doc.add_paragraph(numbered_text, style='List Number')
                
                # Handle tables (simple markdown table format)
                elif '|' in line and not line.startswith('|'):
                    # This is a table row, collect all table rows
                    table_rows = []
                    j = i
                    while j < len(lines) and '|' in lines[j]:
                        if lines[j].strip() and not lines[j].strip().startswith('|---'):
                            table_rows.append(lines[j])
                        j += 1
                    
                    if table_rows:
                        # Create table
                        table_data = []
                        for row in table_rows:
                            cells = [cell.strip() for cell in row.split('|') if cell.strip()]
                            table_data.append(cells)
                        
                        if table_data:
                            table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                            table.style = 'Table Grid'
                            
                            for row_idx, row_data in enumerate(table_data):
                                for col_idx, cell_data in enumerate(row_data):
                                    if col_idx < len(table.rows[row_idx].cells):
                                        table.rows[row_idx].cells[col_idx].text = cell_data
                    
                    i = j - 1
                
                # Handle bold text
                elif '**' in line:
                    paragraph = doc.add_paragraph()
                    parts = line.split('**')
                    for part_idx, part in enumerate(parts):
                        if part_idx % 2 == 0:
                            paragraph.add_run(part)
                        else:
                            run = paragraph.add_run(part)
                            run.bold = True
                
                # Handle regular paragraphs
                else:
                    doc.add_paragraph(line)
                
                i += 1
            
            # Save document
            doc.save(output_path)
            logger.info(f"Word document saved to {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error converting to Word format: {e}")
            raise
    
    def convert_to_text(self, markdown_content: str, output_path: str) -> str:
        """Convert markdown content to plain text."""
        try:
            # Remove markdown formatting
            text_content = markdown_content
            
            # Remove headers
            text_content = re.sub(r'^#+\s*', '', text_content, flags=re.MULTILINE)
            
            # Remove bold formatting
            text_content = re.sub(r'\*\*(.*?)\*\*', r'\1', text_content)
            
            # Remove italic formatting
            text_content = re.sub(r'\*(.*?)\*', r'\1', text_content)
            
            # Remove links
            text_content = re.sub(r'\[([^\]]+)\]\([^\)]+\)', r'\1', text_content)
            
            # Remove code blocks
            text_content = re.sub(r'```.*?```', '', text_content, flags=re.DOTALL)
            
            # Remove inline code
            text_content = re.sub(r'`([^`]+)`', r'\1', text_content)
            
            # Clean up table formatting
            text_content = re.sub(r'\|', ' | ', text_content)
            text_content = re.sub(r'^\s*\|\s*', '', text_content, flags=re.MULTILINE)
            text_content = re.sub(r'\s*\|\s*$', '', text_content, flags=re.MULTILINE)
            
            # Remove table separators
            text_content = re.sub(r'^[-:\s|]+$', '', text_content, flags=re.MULTILINE)
            
            # Clean up extra whitespace
            text_content = re.sub(r'\n\s*\n\s*\n', '\n\n', text_content)
            text_content = text_content.strip()
            
            # Save text file
            with open(output_path, 'w', encoding='utf-8') as f:
                f.write(text_content)
            
            logger.info(f"Text document saved to {output_path}")
            return output_path
            
        except Exception as e:
            logger.error(f"Error converting to text format: {e}")
            raise
    
    def convert_report(self, markdown_path: str, output_format: str) -> str:
        """Convert a markdown report to the specified format."""
        if output_format not in self.supported_formats:
            raise ValueError(f"Unsupported format: {output_format}")
        
        if output_format == 'md':
            return markdown_path
        
        # Read markdown content
        with open(markdown_path, 'r', encoding='utf-8') as f:
            markdown_content = f.read()
        
        # Generate output path
        input_path = Path(markdown_path)
        output_path = input_path.with_suffix(f'.{output_format}')
        
        # Convert based on format
        if output_format == 'docx':
            return self.convert_to_word(markdown_content, str(output_path))
        elif output_format == 'txt':
            return self.convert_to_text(markdown_content, str(output_path))
        
        return str(output_path)
